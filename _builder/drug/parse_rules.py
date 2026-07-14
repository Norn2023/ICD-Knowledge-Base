"""Parse all audit rule Excel files and build a structured rule database"""
import openpyxl, json, sys, os, glob, re
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

RULES_DIR = 'D:/AI_libra/codex_Obsi/_builder/drug'
OUTPUT_PATH = 'D:/AI_libra/codex_Obsi/_builder/drug/audit_rules.json'

def safe_str(v):
    """Convert cell value to string safely"""
    if v is None:
        return ''
    return str(v).strip()

def find_data_start(ws):
    """Find the row where actual column headers and data start"""
    for row_idx in range(1, min(10, ws.max_row + 1)):
        row = [safe_str(ws.cell(row=row_idx, column=c).value) for c in range(1, ws.max_column + 1)]
        # Look for known column headers
        text = ' '.join(row)
        if '序号' in row and ('药品通用名' in row or '药品名称' in row):
            return row_idx
    return 4  # default

def parse_rules():
    all_rules = []

    # Find all xlsx files
    xlsx_files = []
    for root, dirs, files in os.walk(RULES_DIR):
        for f in files:
            if f.endswith('.xlsx') and '汇总' not in f and '药品目录' not in f:
                xlsx_files.append(os.path.join(root, f))

    for fp in sorted(xlsx_files):
        rel_path = os.path.relpath(fp, RULES_DIR)
        dir_name = os.path.basename(os.path.dirname(fp))

        try:
            wb = openpyxl.load_workbook(fp, data_only=True)
            ws = wb.active
            data_start = find_data_start(ws)

            # Get column headers
            headers = [safe_str(ws.cell(row=data_start, column=c).value) for c in range(1, ws.max_column + 1)]
            print(f"\n[{dir_name}] {os.path.basename(fp)}")
            print(f"  Headers: {headers}")

            # Determine rule type from directory name
            rule_type = dir_name

            # Parse data rows
            rules_found = 0
            for row_idx in range(data_start + 1, ws.max_row + 1):
                row = [safe_str(ws.cell(row=row_idx, column=c).value) for c in range(1, ws.max_column + 1)]

                # Skip empty rows
                if not any(row):
                    continue

                # Check if this is a data row (has sequence number or drug name)
                drug_name = ''
                drug_code = ''
                check_logic = ''
                logic_basis = ''
                seq = ''

                for col_idx, header in enumerate(headers):
                    if col_idx >= len(row):
                        break
                    val = row[col_idx]

                    if header == '序号' and val.isdigit():
                        seq = val
                    elif header == '药品通用名':
                        if not drug_name:
                            drug_name = val
                    elif header == '药品名称':
                        if not drug_name:
                            drug_name = val
                    elif header == '中药饮片名称':
                        if not drug_name:
                            drug_name = val
                    elif header == '检出逻辑':
                        check_logic = val
                    elif header == '逻辑依据':
                        logic_basis = val
                    elif '药品代码' in header and len(val) >= 15:
                        drug_code = val
                    elif '限定性别' in header:
                        check_logic = val  # 性别限制规则
                    elif header == '是否为2024年医保目录内\n药品' or header.startswith('是否为'):
                        pass  # metadata

                # For 15th batch and others with '药品代码' column
                if not drug_name:
                    for col_idx, header in enumerate(headers):
                        if col_idx < len(row) and '药品通用名' in header:
                            drug_name = row[col_idx]
                        elif col_idx < len(row) and '药品代码' in header and len(row[col_idx]) >= 15:
                            drug_code = row[col_idx]

                if drug_name or drug_code:
                    rule = {
                        'rule_type': rule_type,
                        'source_file': os.path.basename(fp),
                        'source_dir': dir_name,
                        'drug_name': drug_name,
                        'drug_code': drug_code if len(drug_code) >= 15 else '',
                        'check_logic': check_logic,
                        'logic_basis': logic_basis,
                    }
                    all_rules.append(rule)
                    rules_found += 1

            print(f"  Parsed {rules_found} rules")
            wb.close()

        except Exception as e:
            print(f"  Error [{dir_name}]: {e}")

    print(f"\nTotal rules parsed: {len(all_rules)}")

    # Save
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(all_rules, f, ensure_ascii=False, indent=1)

    # Summary by type
    type_counts = defaultdict(int)
    for r in all_rules:
        type_counts[r['rule_type']] += 1
    print("\nRules by type:")
    for t, c in sorted(type_counts.items()):
        print(f"  {t}: {c}")

    return all_rules

if __name__ == '__main__':
    parse_rules()
