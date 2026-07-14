"""
Rebuild obstetrics mappings using the two manual mapping files:
1. 上海产科类医疗服务价格项目.xlsx - Shanghai obstetrics price items
2. 上海市产科类医疗服务价格项目立项指南映射关系（参考）.xlsx - Manual mapping
"""
import json
import gzip
import os
import pandas as pd
import openpyxl
from difflib import SequenceMatcher

BASE = os.path.dirname(os.path.abspath(__file__))
DATA_FILE = os.path.join(BASE, 'data.json')
BUILDER = os.path.dirname(BASE)

def clean(s):
    return s.replace('\n','').replace(' ','').replace('（','(').replace('）',')').strip()

def fuzzy_match(target, candidates, key='name', threshold=0.55):
    ct = clean(target)
    best_score = 0
    best_idx = -1
    for i, c in enumerate(candidates):
        cn = clean(c.get(key, ''))
        if not cn:
            continue
        if ct == cn:
            return i, 1.0
        if ct in cn or cn in ct:
            return i, 0.95
        score = SequenceMatcher(None, ct, cn).ratio()
        if score > best_score:
            best_score = score
            best_idx = i
    if best_score >= threshold:
        return best_idx, best_score
    return -1, 0

def main():
    # ========================================================
    # 1. Load Shanghai obstetrics price items
    # ========================================================
    sh_ob_file = os.path.join(BUILDER, '上海产科类医疗服务价格项目.xlsx')
    tables = pd.read_html(sh_ob_file, encoding='utf-8')
    df = tables[0]

    # Extract data rows (skip headers at row 0-1, category row 2)
    sh_ob_items = []
    for _, row in df.iterrows():
        seq = str(row.iloc[0]).strip() if pd.notna(row.iloc[0]) else ''
        code = str(row.iloc[1]).strip() if pd.notna(row.iloc[1]) else ''
        name = str(row.iloc[2]).strip() if pd.notna(row.iloc[2]) else ''
        # Skip headers and category rows
        if not code or len(code) < 5:
            continue
        # Skip rows that are just category labels
        if code.isdigit() and len(code) < 6:
            continue
        desc = str(row.iloc[3]).strip() if pd.notna(row.iloc[3]) else ''
        cost = str(row.iloc[4]).strip() if pd.notna(row.iloc[4]) else ''
        unit = str(row.iloc[5]).strip() if pd.notna(row.iloc[5]) else ''
        price = str(row.iloc[6]).strip() if pd.notna(row.iloc[6]) else ''
        notes = str(row.iloc[9]).strip() if pd.notna(row.iloc[9]) else ''

        if name:
            sh_ob_items.append({
                'seq': seq,
                'code': code,
                'name': name,
                'desc': desc[:200],
                'unit': unit,
                'price': price,
                'notes': notes[:200],
            })

    print(f'Loaded {len(sh_ob_items)} Shanghai obstetrics items')

    # ========================================================
    # 2. Load manual mapping reference
    # ========================================================
    map_file = os.path.join(BUILDER, '上海市产科类医疗服务价格项目立项指南映射关系（参考）.xlsx')
    wb_map = openpyxl.load_workbook(map_file, read_only=True, data_only=True)
    ws = wb_map[wb_map.sheetnames[0]]

    manual_mappings = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        seq = str(row[0]).strip() if row[0] else ''
        pg_name = str(row[1]).strip() if row[1] else ''
        pg_surcharge = str(row[2]).strip() if row[2] else ''
        pg_extension = str(row[3]).strip() if row[3] else ''
        sh_names = str(row[4]).strip() if row[4] else ''
        ts_names = str(row[5]).strip() if row[5] else ''

        if seq and pg_name:
            manual_mappings.append({
                'seq': seq,
                'pg_name': pg_name,
                'pg_surcharge': [s.strip() for s in pg_surcharge.split('\n') if s.strip()],
                'pg_extension': [s.strip() for s in pg_extension.split('\n') if s.strip()],
                'sh_names': [s.strip() for s in sh_names.split('\n') if s.strip()],
                'ts_names': [s.strip() for s in ts_names.split('\n') if s.strip()],
            })
    wb_map.close()
    print(f'Loaded {len(manual_mappings)} manual mappings')

    # ========================================================
    # 3. Load current mapping data
    # ========================================================
    with open(DATA_FILE, 'r', encoding='utf-8') as f:
        data = json.load(f)

    sh_all = data['shanghai']
    pg_all = data['pricing_guide']
    ts_all = data['tech_specs']
    cross_refs = data['cross_refs']

    # ========================================================
    # 4. First, add Shanghai obstetrics items if not already present
    # ========================================================
    new_sh_added = 0
    for item in sh_ob_items:
        # Check if already exists by code
        exists = any(s['code'] == item['code'] for s in sh_all)
        if not exists:
            sh_id = f'sh_{len(sh_all)}'
            sh_all.append({
                'id': sh_id,
                'code': item['code'],
                'name': item['name'],
                'desc': item['desc'],
                'unit': item['unit'],
                'price': item['price'],
                'notes': item['notes'],
            })
            new_sh_added += 1

    print(f'Added {new_sh_added} new Shanghai items')

    # ========================================================
    # 5. Remove old obstetrics cross-refs (those involving obstetric SH codes)
    # ========================================================
    ob_sh_codes = set(item['code'] for item in sh_ob_items)
    old_refs_count = len(cross_refs)
    cross_refs = [ref for ref in cross_refs if ref['sh_code'] not in ob_sh_codes]
    removed = old_refs_count - len(cross_refs)
    print(f'Removed {removed} old obstetrics cross-refs')

    # ========================================================
    # 6. Build new cross-refs from manual mapping
    # ========================================================
    new_refs = 0

    for mm in manual_mappings:
        pg_name = mm['pg_name']
        sh_names = mm['sh_names']
        ts_names = mm['ts_names']

        # Find pricing guide item
        pg_idx, pg_score = fuzzy_match(pg_name, pg_all, 'name', threshold=0.6)
        if pg_idx < 0:
            print(f'  [SKIP PG] {pg_name}')
            continue

        pg = pg_all[pg_idx]
        print(f'[OK] {pg_name} -> PG:{pg["name"]}')

        # Find tech spec items
        ts_matches = []
        for ts_name in ts_names:
            ts_idx, ts_score = fuzzy_match(ts_name, ts_all, 'name', threshold=0.55)
            if ts_idx >= 0:
                ts_matches.append({
                    'id': ts_all[ts_idx]['id'],
                    'code': ts_all[ts_idx]['code'],
                    'name': ts_all[ts_idx]['name'],
                    'score': ts_score,
                })

        # For each Shanghai price name in this mapping
        for sh_name in sh_names:
            # Match against our full SH list
            sh_idx, sh_score = fuzzy_match(sh_name, sh_all, 'name', threshold=0.55)
            if sh_idx < 0:
                print(f'  [SKIP SH] {sh_name}')
                continue

            sh = sh_all[sh_idx]
            ts_id = ts_matches[0]['id'] if ts_matches else ''
            ts_code = ts_matches[0]['code'] if ts_matches else ''
            ts_name_match = ts_matches[0]['name'] if ts_matches else ''
            ts_score_str = f'{ts_matches[0]["score"]:.2f}' if ts_matches else ''

            ref = {
                'sh_id': sh['id'],
                'sh_name': sh['name'],
                'sh_code': sh['code'],
                'pg_id': pg['id'],
                'pg_name': pg_name,
                'pg_seq': pg['seq'],
                'pg_category': pg['category'],
                'ts_id': ts_id,
                'ts_name': ts_name_match,
                'ts_code': ts_code,
                'icd9_id': '',
                'icd9_name': '',
                'icd9_code': '',
                'sh_pg_score': '0.95',  # Manual mapping = high confidence
                'pg_ts_score': ts_score_str,
            }
            cross_refs.append(ref)
            new_refs += 1
            print(f'    SH: {sh_name} -> {sh["id"]} (code: {sh["code"]})')

    # ========================================================
    # 7. Rebuild index
    # ========================================================
    sh_index = {}
    pg_index = {}
    ts_index = {}
    icd9_index = {}
    for i, ref in enumerate(cross_refs):
        if ref['sh_id']: sh_index.setdefault(ref['sh_id'], []).append(i)
        if ref['pg_id']: pg_index.setdefault(ref['pg_id'], []).append(i)
        if ref['ts_id']: ts_index.setdefault(ref['ts_id'], []).append(i)
        if ref['icd9_id']: icd9_index.setdefault(ref['icd9_id'], []).append(i)

    data['shanghai'] = sh_all
    data['cross_refs'] = cross_refs
    data['_index'] = {'sh': sh_index, 'pg': pg_index, 'ts': ts_index, 'icd9': icd9_index}
    data['stats'] = {
        'shanghai_count': len(sh_all),
        'pricing_guide_count': len(pg_all),
        'tech_specs_count': len(ts_all),
        'icd9_count': len(data['icd9']),
        'cross_refs_count': len(cross_refs),
        'shanghai_matched': sum(1 for r in cross_refs if r['pg_id']),
        'full_chain': sum(1 for r in cross_refs if r['sh_id'] and r['pg_id'] and r['ts_id'] and r['icd9_id']),
    }

    # ========================================================
    # 8. Save
    # ========================================================
    with open(DATA_FILE, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False)

    json_bytes = json.dumps(data, ensure_ascii=False).encode('utf-8')
    with gzip.open(DATA_FILE + '.gz', 'wb', compresslevel=6) as f:
        f.write(json_bytes)

    print(f'\n=== Results ===')
    print(f'  Shanghai items: {len(sh_all)} (+{new_sh_added})')
    print(f'  Old refs removed: {removed}')
    print(f'  New manual refs: {new_refs}')
    print(f'  Total cross-refs: {len(cross_refs)}')
    stats = data['stats']
    print(f'  Shanghai matched: {stats["shanghai_matched"]}')
    print(f'  Full chain: {stats["full_chain"]}')
    print(f'  File size: {len(json_bytes)} bytes')

if __name__ == '__main__':
    main()
