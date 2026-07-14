"""Fix source badges properly - parse ref file with parent tracking"""
import openpyxl, gzip, json, os

def clean(s):
    if not s: return ''
    return s.replace('\n','').replace(' ','').replace('　','').replace('（','(').replace('）',')').replace(';','').strip()

BASE = '4way-mapping'

# Parse OB reference file
ref_file = None
for f in os.listdir('.'):
    if '产科' in f and '映射' in f and f.endswith('.xlsx'):
        ref_file = f
        break

pg_map = {}  # clean PG name → {'sh': set((name, source)), 'ts': set((name, source))}
last_pg = ''  # track parent PG for surcharge/extension rows

if ref_file:
    wb = openpyxl.load_workbook(ref_file, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows(min_row=4, values_only=True):
        seq = str(row[0]).strip() if row[0] else ''
        pg_name = str(row[1]).strip().replace('\n','') if row[1] else ''
        surcharge = str(row[2]).strip() if len(row) > 2 and row[2] else ''
        extension = str(row[3]).strip() if len(row) > 3 and row[3] else ''
        sh_text = str(row[4]).strip() if len(row) > 4 and row[4] else ''
        ts_text = str(row[5]).strip() if len(row) > 5 and row[5] else ''

        if pg_name and seq:
            last_pg = clean(pg_name)
            pg_map.setdefault(last_pg, {'sh': set(), 'ts': set()})
            for s in sh_text.split('\n'):
                s = s.strip().rstrip(';').rstrip('；')
                if s:
                    pg_map[last_pg]['sh'].add((s, 'project_name'))
            for t in ts_text.split('\n'):
                t = t.strip()
                if t:
                    pg_map[last_pg]['ts'].add((t, 'project_name'))
        elif surcharge or extension:
            # Sub-row: belongs to last_pg, items are 'into_price'
            sub_ts = [t.strip() for t in ts_text.split('\n') if t.strip()] if ts_text else []
            for t in sub_ts:
                pg_map.setdefault(last_pg, {'sh': set(), 'ts': set()})
                pg_map[last_pg]['ts'].add((t, 'into_price'))

    wb.close()
    print(f'OB ref: {len(pg_map)} PG entries')

# Also try rehab
rehab_file = None
for f in os.listdir('.'):
    if '康复' in f and '映射' in f and f.endswith('.xlsx'):
        try:
            with open(f, 'rb') as test:
                if test.read(2) == b'PK':
                    rehab_file = f
        except:
            pass

if rehab_file:
    try:
        wb = openpyxl.load_workbook(rehab_file, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        last_pg = ''
        for row in ws.iter_rows(min_row=4, values_only=True):
            seq = str(row[0]).strip() if row[0] else ''
            pg_name = str(row[1]).strip().replace('\n','') if row[1] else ''
            surcharge = str(row[2]).strip() if len(row) > 2 and row[2] else ''
            sh_text = str(row[4]).strip() if len(row) > 4 and row[4] else ''
            ts_text = str(row[5]).strip() if len(row) > 5 and row[5] else ''
            if pg_name and seq:
                last_pg = clean(pg_name)
                pg_map.setdefault(last_pg, {'sh': set(), 'ts': set()})
                for s in sh_text.split('\n'):
                    s = s.strip().rstrip(';').rstrip('；')
                    if s: pg_map[last_pg]['sh'].add((s, 'project_name'))
                for t in ts_text.split('\n'):
                    t = t.strip()
                    if t: pg_map[last_pg]['ts'].add((t, 'project_name'))
            elif surcharge:
                sub_ts = [t.strip() for t in ts_text.split('\n') if t.strip()] if ts_text else []
                for t in sub_ts:
                    pg_map.setdefault(last_pg, {'sh': set(), 'ts': set()})
                    pg_map[last_pg]['ts'].add((t, 'into_price'))
        wb.close()
        print(f'Rehab ref OK')
    except Exception as e:
        print(f'Rehab skipped: {e}')

# Load current data
with gzip.open(f'{BASE}/data.json.gz', 'rb') as f:
    data = json.loads(f.read())

cross_refs = data['cross_refs']

# Match cross_refs with source info
def match_name(ref_name, ref_set):
    """Match a name against a set of (name, source) tuples."""
    cname = clean(ref_name)
    if not cname:
        return None
    for rname, source in ref_set:
        crname = clean(rname)
        if crname == cname or crname in cname or cname in crname:
            return source
    return None

matched = {'project_name': 0, 'into_price': 0, 'total': 0}
for ref in cross_refs:
    pg_name = ref.get('pg_name', '')
    cname = clean(pg_name)
    if cname not in pg_map:
        continue
    pge = pg_map[cname]

    # Match SH
    sh_name = ref.get('sh_name', '')
    if sh_name and pge['sh']:
        s = match_name(sh_name, pge['sh'])
        if s:
            ref['sh_source'] = s
            matched[s] = matched.get(s, 0) + 1
            matched['total'] += 1

    # Match TS
    ts_name = ref.get('ts_name', '')
    if ts_name and pge['ts']:
        s = match_name(ts_name, pge['ts'])
        if s:
            ref['ts_source'] = s
            matched['total'] += 1

# Default unmapped to project_name
for ref in cross_refs:
    if ref.get('ts_name') and not ref.get('ts_source'):
        ref['ts_source'] = 'project_name'
    if ref.get('sh_name') and not ref.get('sh_source'):
        ref['sh_source'] = 'project_name'

print(f'Source match: {matched}')

# Save
with gzip.open(f'{BASE}/data.json.gz', 'wb', compresslevel=6) as f:
    f.write(json.dumps(data, ensure_ascii=False).encode('utf-8'))
print(f'Saved: {len(cross_refs)} refs')
