"""Fix cross-refs with source badges and correct ICD-9 for 会阴裂伤"""
import openpyxl, gzip, json, os

def clean(s):
    return s.replace('\n','').replace(' ','').replace('　','').replace('（','(').replace('）',')').replace(';','').strip()

BASE = '4way-mapping'

# 1. Parse reference file with source distinction
ref_file = None
for f in os.listdir('.'):
    if '产科' in f and '映射' in f and f.endswith('.xlsx'):
        ref_file = f
        break

pg_map = {}  # PG_name -> {sh: set of (name, source), ts: set of (name, source)}

if ref_file:
    wb = openpyxl.load_workbook(ref_file, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows(min_row=4, values_only=True):
        seq = str(row[0]).strip() if row[0] else ''
        pg_name = str(row[1]).strip().replace('\n','') if row[1] else ''
        surcharge = str(row[2]).strip() if len(row) > 2 and row[2] else ''
        extension = str(row[3]).strip() if len(row) > 3 and row[3] else ''
        sh_text = str(row[4]).strip() if len(row) > 4 and row[4] else ''
        ts_text = str(row[5]).strip() if len(row) > 5 and row[5] else ''

        if pg_name and seq:
            # Main row: project_name
            for s in sh_text.split('\n'):
                s = s.strip().rstrip(';').rstrip('；')
                if s:
                    pg_map.setdefault(clean(pg_name), {'sh': set(), 'ts': set()})
                    pg_map[clean(pg_name)]['sh'].add((s, 'project_name'))
            for t in ts_text.split('\n'):
                t = t.strip()
                if t:
                    pg_map.setdefault(clean(pg_name), {'sh': set(), 'ts': set()})
                    pg_map[clean(pg_name)]['ts'].add((t, 'project_name'))

        if surcharge or extension:
            # Sub-row: into_price for TS items
            sub_ts = [t.strip() for t in ts_text.split('\n') if t.strip()] if ts_text else []
            for t in sub_ts:
                for cname in pg_map:
                    pg_map[cname]['ts'].add((t, 'into_price'))
    wb.close()
    print(f'OB ref parsed: {len(pg_map)} PG entries')

# Also parse rehab reference
rehab_file = None
for f in os.listdir('.'):
    if '康复' in f and '映射' in f and f.endswith('.xlsx'):
        try:
            with open(f, 'rb') as test:
                if test.read(2) == b'PK':
                    rehab_file = f
        except:
            pass

if rehab_file:
    try:
        wb = openpyxl.load_workbook(rehab_file, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        for row in ws.iter_rows(min_row=4, values_only=True):
            seq = str(row[0]).strip() if row[0] else ''
            pg_name = str(row[1]).strip().replace('\n','') if row[1] else ''
            surcharge = str(row[2]).strip() if len(row) > 2 and row[2] else ''
            sh_text = str(row[4]).strip() if len(row) > 4 and row[4] else ''
            ts_text = str(row[5]).strip() if len(row) > 5 and row[5] else ''

            if pg_name and seq:
                for s in sh_text.split('\n'):
                    s = s.strip().rstrip(';').rstrip('；')
                    if s:
                        pg_map.setdefault(clean(pg_name), {'sh': set(), 'ts': set()})
                        pg_map[clean(pg_name)]['sh'].add((s, 'project_name'))
                for t in ts_text.split('\n'):
                    t = t.strip()
                    if t:
                        pg_map.setdefault(clean(pg_name), {'sh': set(), 'ts': set()})
                        pg_map[clean(pg_name)]['ts'].add((t, 'project_name'))

            if surcharge:
                sub_ts = [t.strip() for t in ts_text.split('\n') if t.strip()] if ts_text else []
                for t in sub_ts:
                    for cname in pg_map:
                        pg_map[cname]['ts'].add((t, 'into_price'))
        wb.close()
        print(f'Rehab ref parsed OK')
    except Exception as e:
        print(f'Rehab ref skipped: {e}')

print(f'Total PG entries: {len(pg_map)}')

# 2. Load current data
with gzip.open(f'{BASE}/data.json.gz', 'rb') as f:
    data = json.loads(f.read())

cross_refs = data['cross_refs']
pg_items = data['pricing_guide']

# Build lookups
pg_by_clean = {clean(p['name']): p for p in pg_items}

# 3. Update each cross-ref with source info
updated = 0
for ref in cross_refs:
    pg_name = ref.get('pg_name', '')
    cname = clean(pg_name)
    if cname not in pg_map:
        continue

    sh_name = ref.get('sh_name', '')
    if sh_name and pg_map[cname]['sh']:
        csh = clean(sh_name)
        for ref_name, source in pg_map[cname]['sh']:
            cref = clean(ref_name)
            if cref == csh or cref in csh or csh in cref:
                ref['sh_source'] = source
                break

    ts_name = ref.get('ts_name', '')
    if ts_name and pg_map[cname]['ts']:
        cts = clean(ts_name)
        for ref_name, source in pg_map[cname]['ts']:
            cref = clean(ref_name)
            if cref == cts or cref in cts or cts in cref:
                ref['ts_source'] = source
                break

    updated += 1

print(f'Updated {updated} cross-refs with source badges')

# 4. Fix ICD-9 for 会阴裂伤
icd9_fix_code = '75.6902'
icd9_fix_name = '近期产科会阴裂伤修补术'
icd9_id = ''
for icd in data['icd9']:
    if icd.get('yb_code') == icd9_fix_code:
        icd9_id = icd['id']
        break

fixed_icd9 = 0
for ref in cross_refs:
    tsn = ref.get('ts_name', '')
    if '会阴' in tsn and '裂伤' in tsn:
        ref['icd9_code'] = icd9_fix_code
        ref['icd9_name'] = icd9_fix_name
        ref['icd9_id'] = icd9_id
        fixed_icd9 += 1

print(f'Fixed ICD-9: {fixed_icd9} 会阴裂伤 refs')

# 5. Rebuild index
sh_idx = {}; pg_idx = {}; ts_idx = {}; icd9_idx = {}
for i, ref in enumerate(cross_refs):
    if ref['sh_id']: sh_idx.setdefault(ref['sh_id'], []).append(i)
    if ref['pg_id']: pg_idx.setdefault(ref['pg_id'], []).append(i)
    if ref['ts_id']: ts_idx.setdefault(ref['ts_id'], []).append(i)
    if ref['icd9_id']: icd9_idx.setdefault(ref['icd9_id'], []).append(i)

data['_index'] = {'sh': sh_idx, 'pg': pg_idx, 'ts': ts_idx, 'icd9': icd9_idx}

dc = data.setdefault('stats', {})
dc['cross_refs_count'] = len(cross_refs)
dc['full_chain'] = sum(1 for r in cross_refs if r['sh_id'] and r['pg_id'] and r['ts_id'] and r['icd9_id'])

# Save
with open(f'{BASE}/data.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False)
jbytes = json.dumps(data, ensure_ascii=False).encode('utf-8')
with gzip.open(f'{BASE}/data.json.gz', 'wb', compresslevel=6) as f:
    f.write(jbytes)

print(f'\nFinal: {len(cross_refs)} refs, {dc["full_chain"]} full chain')
print(f'SH with source: {sum(1 for r in cross_refs if r.get("sh_source"))}')
print(f'TS with source: {sum(1 for r in cross_refs if r.get("ts_source"))}')
